import os
from pptx import Presentation

path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3.pptx'
prs = Presentation(path)

# Slide 1: Intro / Accroche (English)
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Alice" in text or "Ava" in text or "Hi" in text:
            shape.text_frame.text = "Hi, I'm Alice"
        elif "guide" in text.lower() or "journey" in text.lower():
            shape.text_frame.text = "Hello, I'm Alice. Today, I want to address the most valuable subject of your next meeting: attention. In a world where everyone is talking, the true challenge is being heard. At Zalyus Slide, we make sure your voice stands out."
        elif "Onboarding Training" in text:
            shape.text_frame.text = "Zalyus Slide"

slide1.notes_slide.notes_text_frame.text = "Alice: « Hello, I'm Alice. Today, I want to address the most valuable subject of your next meeting: attention. In a world where everyone is talking, the true challenge is being heard. At Zalyus Slide, we make sure your voice stands out. »"


# Slide 2: The Stakes / Perception (English)
slide2 = prs.slides[1]
for shape in slide2.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Zalyus" in text or "Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "mission" in text.lower() or "teams" in text.lower():
            shape.text_frame.text = "We've all seen it: brilliant ideas drowned under messy, amateur slides. But here is the truth: your audience judges your expertise through your visuals. At Zalyus Slide, we believe your presentation shouldn't just be a background; it must be your ultimate proof of authority."

slide2.notes_slide.notes_text_frame.text = "Alice: « We've all seen it: brilliant ideas drowned under messy, amateur slides. But here is the truth: your audience judges your expertise through your visuals. At Zalyus Slide, we believe your presentation shouldn't just be a background; it must be your ultimate proof of authority. »"


# Slide 3: Discover Zalyus Slide / The Bridge (English)
slide3 = prs.slides[2]
for shape in slide3.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Zalyus" in text or "Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "values" in text.lower() or "One Skill" in text.lower():
            shape.text_frame.text = "That is why we created Zalyus Slide. We are much more than simple designers; we are visual strategists. Our mission is simple: bridge the gap between your strategic vision and the professional excellence your clients expect."

slide3.notes_slide.notes_text_frame.text = "Alice: « That is why we created Zalyus Slide. We are much more than simple designers; we are visual strategists. Our mission is simple: bridge the gap between your strategic vision and the professional excellence your clients expect. »"


# Slide 4: Visual Engineering / Methodology (English)
slide4 = prs.slides[3]
for shape in slide4.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Zalyus" in text or "Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "culture" in text.lower() or "teamwork" in text.lower():
            shape.text_frame.text = "How do we do it? We call it visual engineering. We don't just 'pretty things up.' We use cognitive psychology and data storytelling to restructure your message, making complex information clear, fluid, and impossible to ignore."

slide4.notes_slide.notes_text_frame.text = "Alice: « How do we do it? We call it visual engineering. We don't just 'pretty things up.' We use cognitive psychology and data storytelling to restructure your message, making complex information clear, fluid, and impossible to ignore. »"


# Slide 5: Capabilities / Tools (English)
slide5 = prs.slides[4]
for shape in slide5.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "First Week" in text:
            shape.text_frame.text = "Capabilities & Ecosystem"
        elif "first week" in text.lower() or "meet your team" in text.lower():
            shape.text_frame.text = "From high-stakes investor pitch decks to enterprise document templates, our team masters the entire ecosystem. Whether it's PowerPoint, Keynote, or Canva, we craft custom, high-performing assets that work flawlessly across all platforms."

slide5.notes_slide.notes_text_frame.text = "Alice: « From high-stakes investor pitch decks to enterprise document templates, our team masters the entire ecosystem. Whether it's PowerPoint, Keynote, or Canva, we craft custom, high-performing assets that work flawlessly across all platforms. »"


# Slide 6: The Promise / Reliability (English)
slide6 = prs.slides[5]
for shape in slide6.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Tools & Platforms" in text:
            shape.text_frame.text = "Our Promise"
        elif "messaging" in text.lower() or "connected" in text.lower():
            shape.text_frame.text = "Choosing Zalyus Slide means choosing peace of mind. We deliver agency-grade execution combined with the agility of a dedicated partner. We guarantee total confidentiality, strict adherence to deadlines, and flawless finish every time."

slide6.notes_slide.notes_text_frame.text = "Alice: « Choosing Zalyus Slide means choosing peace of mind. We deliver agency-grade execution combined with the agility of a dedicated partner. We guarantee total confidentiality, strict adherence to deadlines, and flawless finish every time. »"


# Slide 7: Call to Action / Invitation (English)
slide7 = prs.slides[6]
for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Hola" in text or "Alice" in text or "Ava" in text:
            shape.text_frame.text = "I'm Alice"
        elif "guiaré" in text.lower() or "comenzar" in text.lower() or "recorrido" in text.lower() or "guicaré" in text.lower():
            shape.text_frame.text = "I'm Alice, and we are ready to elevate your brand. Don't let an ordinary presentation hold back your potential. Find Zalyus Slide on Upwork, Malt, or Fiverr today. Let's make your next big project unforgettable."
        elif "Capacitación" in text:
            shape.text_frame.text = "Zalyus Slide"

slide7.notes_slide.notes_text_frame.text = "Alice: « I'm Alice, and we are ready to elevate your brand. Don't let an ordinary presentation hold back your potential. Find Zalyus Slide on Upwork, Malt, or Fiverr today. Let's make your next big project unforgettable. »"


# Save output
output_path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3_Alice_Zalyus_English.pptx'
prs.save(output_path)
print(f"Updated English presentation saved to {output_path}")

# Also overwrite the main target file so it has the English script
target_main = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3.pptx'
prs.save(target_main)
print(f"Overwrote source file with English script: {target_main}")
