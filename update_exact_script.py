import os
from pptx import Presentation

path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3.pptx'
prs = Presentation(path)

# Script exact fourni par l'utilisateur
s1 = "Bonjour, je suis Alice. Aujourd’hui, je souhaite aborder le sujet le plus précieux de votre prochaine réunion : l’attention. Dans un monde où tout le monde parle, le véritable défi est de se faire entendre. Chez Zalyus Slide , nous veillons à ce que votre voix se démarque."
s2 = "On l’a tous constaté : des idées brillantes noyées sous des diapositives brouillonnes et amateurs. Mais voici la vérité : votre public juge votre expertise à travers vos visuels. Chez Zalyus Slide, nous pensons que votre présentation ne doit pas se limiter à un simple arrière-plan ; elle doit être votre meilleure preuve d’autorité."
s3 = "C’est pourquoi nous avons créé Zalyus Slide . Nous sommes bien plus que de simples designers ; nous sommes des stratèges visuels. Notre mission est simple : combler le fossé entre votre vision stratégique et l’excellence professionnelle que vos clients attendent."
s4 = "Comment procédons-nous ? Nous appelons cela l’ingénierie visuelle. Nous ne nous contentons pas d’« embellir les choses ». Nous utilisons la psychologie cognitive et la narration de données pour restructurer votre message, rendant ainsi les informations complexes claires, fluides et impossibles à ignorer."
s5 = "Des présentations destinées aux investisseurs les plus importants aux modèles de documents d’entreprise, notre équipe maîtrise l’ensemble de l’écosystème. Qu’il s’agisse de PowerPoint, de Keynote ou de Canva, nous créons des supports personnalisés et performants qui fonctionnent parfaitement sur toutes les plateformes."
s6 = "Choisir Zalyus Slide, c’est choisir la tranquillité d’esprit. Nous offrons une exécution digne d’une agence, alliée à l’agilité d’un partenaire dédié. Nous garantissons une confidentialité totale, le respect strict des délais et une finition impeccable, à chaque fois."
s7 = "Je suis Alice, et nous sommes prêts à sublimer votre marque. Ne laissez pas une présentation ordinaire brider votre potentiel. Retrouvez Zalyus Slide sur Upwork, Malt ou Fiverr dès aujourd’hui. Faisons de votre prochain grand projet un moment inoubliable."

scripts = [s1, s2, s3, s4, s5, s6, s7]

for i, slide in enumerate(prs.slides):
    script_text = scripts[i]

    # Update text shapes on the slide to match or reflect the script & title
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if i == 0:
                if "Ava" in text or "Hi" in text or "Alice" in text:
                    shape.text_frame.text = "Bonjour, je suis Alice"
                elif "guide" in text.lower() or "journey" in text.lower() or "Onboarding" in text:
                    shape.text_frame.text = script_text
                elif "Training" in text:
                    shape.text_frame.text = "Zalyus Slide"
            elif i in [1, 2, 3]:
                if "Our Company" in text:
                    shape.text_frame.text = "Zalyus Slide"
                elif "mission" in text.lower() or "values" in text.lower() or "culture" in text.lower() or "teams" in text.lower() or "At One Skill" in text:
                    shape.text_frame.text = script_text
            elif i == 4:
                if "First Week" in text:
                    shape.text_frame.text = "Capacités"
                elif "first week" in text.lower() or "meet your team" in text.lower() or "During" in text:
                    shape.text_frame.text = script_text
            elif i == 5:
                if "Tools & Platforms" in text:
                    shape.text_frame.text = "La Promesse"
                elif "messaging" in text.lower() or "connected" in text.lower() or "We’ll help" in text:
                    shape.text_frame.text = script_text
            elif i == 6:
                if "Hola" in text or "Alice" in text or "Ava" in text:
                    shape.text_frame.text = "Je suis Alice"
                elif "guiaré" in text.lower() or "comenzar" in text.lower() or "recorrido" in text.lower() or "guicaré" in text.lower() or "Te guiaré" in text:
                    shape.text_frame.text = script_text
                elif "Capacitación" in text:
                    shape.text_frame.text = "Zalyus Slide"

    # Set exact speaker notes for ElevenLabs / Avatar narration
    slide.notes_slide.notes_text_frame.text = f"Alice : « {script_text} »"

# Save presentation to a new filename (to avoid file lock if PowerPoint has it open)
output_path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3_Alice_Zalyus_Exact.pptx'
prs.save(output_path)
print(f"Successfully updated presentation with exact script saved to: {output_path}")
