import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_zalyus_deck():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(10, 25, 47)      # #0A192F Deep Navy
    BG_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Off White
    TEXT_DARK = RGBColor(15, 23, 42)    # #0F172A Slate 900
    TEXT_LIGHT = RGBColor(255, 255, 255)# White
    ACCENT_BLUE = RGBColor(37, 99, 235) # #2563EB Blue 600
    ACCENT_CYAN = RGBColor(14, 165, 233)# #0EA5E9 Sky 500
    CARD_BG = RGBColor(255, 255, 255)   # Pure White
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 Slate 200
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748B Slate 500

    def set_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes_text

    # ==================== DIAPO 1 : L'accroche (Intro d'Alice) ====================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide1, BG_DARK)

    txBox = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.933), Inches(4.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ZALYUS SLIDE"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.name = "Segoe UI"

    p2 = tf.add_paragraph()
    p2.text = "Maîtriser l'attention dans un monde saturé"
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(10)

    p3 = tf.add_paragraph()
    p3.text = "Présentation narrative avec Alice • Consultante & Stratège Visuel"
    p3.font.size = Pt(16)
    p3.font.color.rgb = MUTED_TEXT
    p3.font.name = "Segoe UI"
    p3.space_before = Pt(25)

    add_speaker_notes(slide1,
        "Alice : « Bonjour, je suis Alice. Aujourd’hui, je souhaite aborder le sujet le plus précieux de votre prochaine réunion : l’attention. Dans un monde où tout le monde parle, le véritable défi est de se faire entendre. Chez Zalyus Slide , nous veillons à ce que votre voix se démarque. »"
    )

    # ==================== DIAPO 2 : Les enjeux (perception) ====================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide2, BG_LIGHT)

    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "01. Les Enjeux : La Perception"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Votre public juge votre expertise à travers vos visuels"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub_before = p_sub.space_before = Pt(5)

    # Card
    shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1.5)

    tf_card = shape.text_frame
    tf_card.word_wrap = True
    tf_card.margin_top = Inches(0.5)
    tf_card.margin_left = Inches(0.6)
    tf_card.margin_right = Inches(0.6)

    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Le piège des présentations brouillonnes"
    p_c1.font.size = Pt(22)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_BLUE

    p_c2 = tf_card.add_paragraph()
    p_c2.text = "• Des idées brillantes souvent noyées sous des diapositives amateurs et surchargées.\n• La vérité implacable : votre auditoire évalue la qualité de votre travail et votre crédibilité par la clarté de vos visuels.\n• Votre présentation ne doit pas être un simple arrière-plan : c’est votre meilleure preuve d’autorité."
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(20)

    add_speaker_notes(slide2,
        "Alice : « On l’a tous constaté : des idées brillantes noyées sous des diapositives brouillonnes et amateurs. Mais voici la vérité : votre public juge votre expertise à travers vos visuels. Chez Zalyus Slide, nous pensons que votre présentation ne doit pas se limiter à un simple arrière-plan ; elle doit être votre meilleure preuve d’autorité. »"
    )

    # ==================== DIAPO 3 : Découvrez Zalyus Slide (Le Pont) ====================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide3, BG_LIGHT)

    tb = slide3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "02. Zalyus Slide : Stratèges Visuels"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Combler le fossé entre stratégie et excellence professionnelle"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1.5)

    tf_card = shape.text_frame
    tf_card.word_wrap = True
    tf_card.margin_top = Inches(0.5)
    tf_card.margin_left = Inches(0.6)
    tf_card.margin_right = Inches(0.6)

    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Bien plus que des designers"
    p_c1.font.size = Pt(22)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_CYAN

    p_c2 = tf_card.add_paragraph()
    p_c2.text = "• Une alliance unique entre vision stratégique d'entreprise et impact visuel haut de gamme.\n• Notre mission fondamentale : transformer vos concepts complexes en supports percutants.\n• Apporter aux équipes l'excellence professionnelle et la distinction que vos clients et partenaires exigent."
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(20)

    add_speaker_notes(slide3,
        "Alice : « C’est pourquoi nous avons créé Zalyus Slide . Nous sommes bien plus que de simples designers ; nous sommes des stratèges visuels. Notre mission est simple : combler le fossé entre votre vision stratégique et l’excellence professionnelle que vos clients attendent. »"
    )

    # ==================== DIAPO 4 : Ingénierie visuelle (La méthodologie) ====================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide4, BG_LIGHT)

    tb = slide4.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "03. L'Ingénierie Visuelle"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Psychologie cognitive et narration de données"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1.5)

    tf_card = shape.text_frame
    tf_card.word_wrap = True
    tf_card.margin_top = Inches(0.5)
    tf_card.margin_left = Inches(0.6)
    tf_card.margin_right = Inches(0.6)

    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Au-delà de l'embellissement"
    p_c1.font.size = Pt(22)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_BLUE

    p_c2 = tf_card.add_paragraph()
    p_c2.text = "• Restructuration intelligente de votre message en s'appuyant sur les sciences cognitives.\n• Utilisation avancée du storytelling de données (data-storytelling).\n• Rendre chaque information complexe limpide, fluide et impossible à ignorer par votre auditoire."
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(20)

    add_speaker_notes(slide4,
        "Alice : « Comment procédons-nous ? Nous appelons cela l’ingénierie visuelle. Nous ne nous contentons pas d’« embellir les choses ». Nous utilisons la psychologie cognitive et la narration de données pour restructurer votre message, rendant ainsi les informations complexes claires, fluides et impossibles à ignorer. »"
    )

    # ==================== DIAPO 5 : Capacités (Les outils) ====================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide5, BG_LIGHT)

    tb = slide5.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "04. Écosystème & Capacités"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "PowerPoint, Keynote, Canva : une maîtrise totale multi-plateforme"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    shape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1.5)

    tf_card = shape.text_frame
    tf_card.word_wrap = True
    tf_card.margin_top = Inches(0.5)
    tf_card.margin_left = Inches(0.6)
    tf_card.margin_right = Inches(0.6)

    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Du pitch investisseur au document d'entreprise"
    p_c1.font.size = Pt(22)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_CYAN

    p_c2 = tf_card.add_paragraph()
    p_c2.text = "• Des présentations stratégiques pour les investisseurs de premier plan.\n• Modèles de documents d'entreprise et templates modulaires sur-mesure.\n• Création et optimisation expertes sur PowerPoint, Keynote et Canva avec un rendu impeccable sur tous les supports."
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(20)

    add_speaker_notes(slide5,
        "Alice : « Des présentations destinées aux investisseurs les plus importants aux modèles de documents d’entreprise, notre équipe maîtrise l’ensemble de l’écosystème. Qu’il s’agisse de PowerPoint, de Keynote ou de Canva, nous créons des supports personnalisés et performants qui fonctionnent parfaitement sur toutes les plateformes. »"
    )

    # ==================== DIAPO 6 : La promesse (fiabilité) ====================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide6, BG_LIGHT)

    tb = slide6.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "05. Notre Promesse"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Exécution digne d'une grande agence et agilité d'un partenaire"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1.5)

    tf_card = shape.text_frame
    tf_card.word_wrap = True
    tf_card.margin_top = Inches(0.5)
    tf_card.margin_left = Inches(0.6)
    tf_card.margin_right = Inches(0.6)

    p_c1 = tf_card.paragraphs[0]
    p_c1.text = "Tranquillité d'esprit absolue"
    p_c1.font.size = Pt(22)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_BLUE

    p_c2 = tf_card.add_paragraph()
    p_c2.text = "• Confidentialité totale garantie pour l'ensemble de vos données stratégiques.\n• Respect rigoureux et systématique des délais de livraison.\n• Finition impeccable et exigence de qualité supérieure à chaque projet."
    p_c2.font.size = Pt(16)
    p_c2.font.color.rgb = TEXT_DARK
    p_c2.space_before = Pt(20)

    add_speaker_notes(slide6,
        "Alice : « Choisir Zalyus Slide, c’est choisir la tranquillité d’esprit. Nous offrons une exécution digne d’une agence, alliée à l’agilité d’un partenaire dédié. Nous garantissons une confidentialité totale, le respect strict des délais et une finition impeccable, à chaque fois. »"
    )

    # ==================== DIAPO 7 : Appel à l’action (L’invitation) ====================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide7, BG_DARK)

    txBox = slide7.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Sublimez Votre Marque avec Zalyus Slide"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "Ne laissez pas une présentation ordinaire brider votre potentiel."
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_before = Pt(15)

    p3 = tf.add_paragraph()
    p3.text = "Retrouvez Zalyus Slide dès aujourd'hui sur Upwork, Malt ou Fiverr.\nFaisons de votre prochain grand projet un moment inoubliable."
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_LIGHT
    p3.space_before = Pt(20)

    p4 = tf.add_paragraph()
    p4.text = "— Alice, pour Zalyus Slide"
    p4.font.size = Pt(14)
    p4.font.color.rgb = MUTED_TEXT
    p4.space_before = Pt(30)

    add_speaker_notes(slide7,
        "Alice : « Je suis Alice, et nous sommes prêts à sublimer votre marque. Ne laissez pas une présentation ordinaire brider votre potentiel. Retrouvez Zalyus Slide sur Upwork, Malt ou Fiverr dès aujourd’hui. Faisons de votre prochain grand projet un moment inoubliable. »"
    )

    # Save output locally in repo and target directory
    output_filename = "Zalyus_Slide_Alice_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully saved to {output_filename}")

    target_dir = r"C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs"
    if os.path.exists(target_dir):
        target_path = os.path.join(target_dir, output_filename)
        prs.save(target_path)
        print(f"Presentation also successfully copied to {target_path}")

if __name__ == "__main__":
    create_zalyus_deck()
