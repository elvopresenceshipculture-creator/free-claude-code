import os
from pptx import Presentation

path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3.pptx'
prs = Presentation(path)

# Slide 1: Intro / Accroche
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Ava" in text or "Hi" in text:
            shape.text_frame.text = "Hi, I'm Alice"
        elif "I’ll guide" in text or "journey" in text or "guide" in text.lower():
            shape.text_frame.text = "Bonjour, je suis Alice. Aujourd’hui, je souhaite aborder le sujet le plus précieux de votre prochaine réunion : l’attention. Dans un monde où tout le monde parle, le véritable défi est de se faire entendre. Chez Zalyus Slide , nous veillons à ce que votre voix se démarque."
        elif "Onboarding Training" in text:
            shape.text_frame.text = "Zalyus Slide"

slide1.notes_slide.notes_text_frame.text = "Alice : « Bonjour, je suis Alice. Aujourd’hui, je souhaite aborder le sujet le plus précieux de votre prochaine réunion : l’attention. Dans un monde où tout le monde parle, le véritable défi est de se faire entendre. Chez Zalyus Slide , nous veillons à ce que votre voix se démarque. »"


# Slide 2: Les enjeux (perception)
slide2 = prs.slides[1]
for shape in slide2.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Our Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "Our mission is to help teams" in text or "mission" in text.lower():
            shape.text_frame.text = "On l’a tous constaté : des idées brillantes noyées sous des diapositives brouillonnes et amateurs. Mais voici la vérité : votre public juge votre expertise à travers vos visuels. Chez Zalyus Slide, nous pensons que votre présentation ne doit pas se limiter à un simple arrière-plan ; elle doit être votre meilleure preuve d’autorité."

slide2.notes_slide.notes_text_frame.text = "Alice : « On l’a tous constaté : des idées brillantes noyées sous des diapositives brouillonnes et amateurs. Mais voici la vérité : votre public juge votre expertise à travers vos visuels. Chez Zalyus Slide, nous pensons que votre présentation ne doit pas se limiter à un simple arrière-plan ; elle doit être votre meilleure preuve d’autorité. »"


# Slide 3: Découvrez Zalyus Slide (Le Pont)
slide3 = prs.slides[2]
for shape in slide3.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Our Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "At One Skill Corp. our values" in text or "values" in text.lower():
            shape.text_frame.text = "C’est pourquoi nous avons créé Zalyus Slide . Nous sommes bien plus que de simples designers ; nous sommes des stratèges visuels. Notre mission est simple : combler le fossé entre votre vision stratégique et l’excellence professionnelle que vos clients attendent."

slide3.notes_slide.notes_text_frame.text = "Alice : « C’est pourquoi nous avons créé Zalyus Slide . Nous sommes bien plus que de simples designers ; nous sommes des stratèges visuels. Notre mission est simple : combler le fossé entre votre vision stratégique et l’excellence professionnelle que vos clients attendent. »"


# Slide 4: Ingénierie visuelle (La méthodologie)
slide4 = prs.slides[3]
for shape in slide4.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Our Company" in text:
            shape.text_frame.text = "Zalyus Slide"
        elif "At One Skill Corp. our culture" in text or "culture" in text.lower():
            shape.text_frame.text = "Comment procédons-nous ? Nous appelons cela l’ingénierie visuelle. Nous ne nous contentons pas d’« embellir les choses ». Nous utilisons la psychologie cognitive et la narration de données pour restructurer votre message, rendant ainsi les informations complexes claires, fluides et impossibles à ignorer."

slide4.notes_slide.notes_text_frame.text = "Alice : « Comment procédons-nous ? Nous appelons cela l’ingénierie visuelle. Nous ne nous contentons pas d’« embellir les choses ». Nous utilisons la psychologie cognitive et la narration de données pour restructurer votre message, rendant ainsi les informations complexes claires, fluides et impossibles à ignorer. »"


# Slide 5: Capacités (Les outils)
slide5 = prs.slides[4]
for shape in slide5.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Your First Week" in text:
            shape.text_frame.text = "Écosystème & Capacités"
        elif "During your first week" in text or "meet your team" in text.lower():
            shape.text_frame.text = "Des présentations destinées aux investisseurs les plus importants aux modèles de documents d’entreprise, notre équipe maîtrise l’ensemble de l’écosystème. Qu’il s’agisse de PowerPoint, de Keynote ou de Canva, nous créons des supports personnalisés et performants qui fonctionnent parfaitement sur toutes les plateformes."

slide5.notes_slide.notes_text_frame.text = "Alice : « Des présentations destinées aux investisseurs les plus importants aux modèles de documents d’entreprise, notre équipe maîtrise l’ensemble de l’écosystème. Qu’il s’agisse de PowerPoint, de Keynote ou de Canva, nous créons des supports personnalisés et performants qui fonctionnent parfaitement sur toutes les plateformes. »"


# Slide 6: La promesse (fiabilité)
slide6 = prs.slides[5]
for shape in slide6.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Tools & Platforms" in text:
            shape.text_frame.text = "Notre Promesse"
        elif "We’ll help you stay connected" in text or "messaging" in text.lower():
            shape.text_frame.text = "Choisir Zalyus Slide, c’est choisir la tranquillité d’esprit. Nous offrons une exécution digne d’une agence, alliée à l’agilité d’un partenaire dédié. Nous garantissons une confidentialité totale, le respect strict des délais et une finition impeccable, à chaque fois."

slide6.notes_slide.notes_text_frame.text = "Alice : « Choisir Zalyus Slide, c’est choisir la tranquillité d’esprit. Nous offrons une exécution digne d’une agence, alliée à l’agilité d’un partenaire dédié. Nous garantissons une confidentialité totale, le respect strict des délais et une finition impeccable, à chaque fois. »"


# Slide 7: Appel à l’action (L’invitation)
slide7 = prs.slides[6]
for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if "Hola, soy Ava" in text:
            shape.text_frame.text = "Je suis Alice"
        elif "Te guiaré" in text or "comenzar" in text.lower() or "incorporación" in text.lower():
            shape.text_frame.text = "Je suis Alice, et nous sommes prêts à sublimer votre marque. Ne laissez pas une présentation ordinaire brider votre potentiel. Retrouvez Zalyus Slide sur Upwork, Malt ou Fiverr dès aujourd’hui. Faisons de votre prochain grand projet un moment inoubliable."
        elif "Capacitación" in text:
            shape.text_frame.text = "Zalyus Slide"

slide7.notes_slide.notes_text_frame.text = "Alice : « Je suis Alice, et nous sommes prêts à sublimer votre marque. Ne laissez pas une présentation ordinaire brider votre potentiel. Retrouvez Zalyus Slide sur Upwork, Malt ou Fiverr dès aujourd’hui. Faisons de votre prochain grand projet un moment inoubliable. »"


# Save output
output_path = r'C:\Users\HP\Downloads\Mes présentation powerpoint\Realistic AI Avatars in PowerPoint with ElevenLabs\Interactive Onboarding Training v3_Alice_Zalyus.pptx'
prs.save(output_path)
print(f"Updated presentation saved to {output_path}")

# Also save local copy
prs.save('Interactive Onboarding Training v3_Alice_Zalyus.pptx')
