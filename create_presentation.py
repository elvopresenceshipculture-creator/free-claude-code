import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(10, 25, 47)      # #0A192F Deep Navy
    BG_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Off White
    TEXT_DARK = RGBColor(15, 23, 42)    # #0F172A Slate 900
    TEXT_LIGHT = RGBColor(255, 255, 255)# White
    ACCENT_BLUE = RGBColor(37, 99, 235) # #2563EB Blue 600
    ACCENT_CYAN = RGBColor(14, 165, 233)# #0EA5E9 Sky 500
    CARD_BG = RGBColor(255, 255, 255)   # Pure White
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 Slate 200
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748B Slate 500

    def set_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # ==================== SLIDE 1: TITLE SLIDE ====================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide1, BG_DARK)

    # Title & Subtitle container
    txBox = slide1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "L'ÈRE DES AGENTS AUTONOMES"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.LEFT

    p2 = tf.add_paragraph()
    p2.text = "Redéfinir le futur du logiciel, de l'entreprise et de l'ingénierie"
    p2.font.size = Pt(22)
    p2.font.color.rgb = ACCENT_CYAN
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(15)

    p3 = tf.add_paragraph()
    p3.text = "Présentation interactive générée par Claude • 2026"
    p3.font.size = Pt(14)
    p3.font.color.rgb = MUTED_TEXT
    p3.font.name = "Segoe UI"
    p3.space_before = Pt(40)

    # ==================== SLIDE 2: AGENDA ====================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide2, BG_LIGHT)

    # Header
    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Sommaire de la présentation"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Les 4 axes stratégiques de la transformation autonome"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    # 4 Agenda cards
    cards = [
        ("01", "Mutation des Modèles", "Du chatbot conversationnel à l'exécution de workflows autonomes multi-étapes."),
        ("02", "Architecture & Sécurité", "Fiabilité, vérification rigoureuse et gouvernance des agents en production."),
        ("03", "Cas d'Usage Métier", "Automatisation logicielle, support intelligent et accélération de la R&D."),
        ("04", "Vision & Perspectives", "Vers une symbiose homme-machine durable, éthique et souveraine.")
    ]

    card_width = Inches(2.6)
    card_height = Inches(4.2)
    card_y = Inches(2.2)
    start_x = Inches(1.0)
    gap = Inches(0.31)

    for i, (num, title, desc) in enumerate(cards):
        cx = start_x + i * (card_width + gap)
        # Background card shape
        shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)

        # Text inside card
        tf_card = shape.text_frame
        tf_card.word_wrap = True
        tf_card.margin_top = Inches(0.4)
        tf_card.margin_left = Inches(0.3)
        tf_card.margin_right = Inches(0.3)

        p_num = tf_card.paragraphs[0]
        p_num.text = num
        p_num.font.size = Pt(28)
        p_num.font.bold = True
        p_num.font.color.rgb = ACCENT_BLUE

        p_title = tf_card.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.space_before = Pt(15)

        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = MUTED_TEXT
        p_desc.space_before = Pt(10)

    # ==================== SLIDE 3: CONTENT 2 COLUMNS ====================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide3, BG_LIGHT)

    tb = slide3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "01. La Mutation des Modèles d'IA"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tf.add_paragraph()
    p_sub.text = "Comment l'industrie passe de la simple génération de texte à l'orchestration d'agents"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(5)

    # Two big comparison / feature cards
    col_width = Inches(5.5)
    col_height = Inches(4.5)
    col_y = Inches(2.2)

    cols_data = [
        ("Modèle Réactif Traditionnel", [
            "• Interaction basée sur un tour unique (Question / Réponse).",
            "• Absence de persistance d'état et de mémoire contextuelle complexe.",
            "• Intervention humaine requise à chaque étape de validation.",
            "• Adapté aux tâches textuelles simples et à la recherche documentaire."
        ], ACCENT_BLUE),
        ("Système d'Agents Autonomes", [
            "• Planification dynamique et décomposition de tâches complexes.",
            "• Utilisation d'outils spécialisés (recherche, exécution de code, APIs).",
            "• Boucles de rétroaction et auto-correction en cas d'erreur.",
            "• Capacité d'exécuter des workflows complets de bout en bout."
        ], ACCENT_CYAN)
    ]

    for i, (col_title, items, accent_color) in enumerate(cols_data):
        cx = Inches(1.0) + i * (col_width + Inches(0.333))
        shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, col_y, col_width, col_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)

        tf_col = shape.text_frame
        tf_col.word_wrap = True
        tf_col.margin_top = Inches(0.4)
        tf_col.margin_left = Inches(0.4)
        tf_col.margin_right = Inches(0.4)

        p_t = tf_col.paragraphs[0]
        p_t.text = col_title
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = accent_color

        for item in items:
            p_item = tf_col.add_paragraph()
            p_item.text = item
            p_item.font.size = Pt(15)
            p_item.font.color.rgb = TEXT_DARK
            p_item.space_before = Pt(15)

    # ==================== SLIDE 4: CONCLUSION ====================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide4, BG_DARK)

    txBox = slide4.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Prêt à transformer votre organisation ?"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "L'intégration d'agents autonomes et d'assistants intelligents propulsés par les modèles de pointe n'est plus une option futuriste, c'est le standard de compétitivité de demain."
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_before = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "• Qualité zero-defect • Architecture modulaire • Efficacité maximale"
    p3.font.size = Pt(16)
    p3.font.color.rgb = MUTED_TEXT
    p3.space_before = Pt(30)

    output_path = "presentation_exemple.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
